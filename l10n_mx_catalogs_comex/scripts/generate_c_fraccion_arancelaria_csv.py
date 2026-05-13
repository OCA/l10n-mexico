#!/usr/bin/python3

####
# Fetch tariff codes (fracción arancelaria) with NICO from catalogs posted at:
# https://www.snice.gob.mx/cs/avi/snice/nico.ligie.html
# To generate the data file for the Odoo module l10n_mx_catalogs.c_fraccion
##

import argparse
import csv
import io
import re
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

SOURCE_URL = "https://www.snice.gob.mx/cs/avi/snice/nico.ligie.html"
DEFAULT_OUTPUT = "../data/l10n_mx_catalogs.c_fraccion.csv"
FULL_FRACCION_RE = re.compile(r"^\d{4}\.\d{2}\.\d{2}$")


def clean_text(value):
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return " ".join(str(value).strip().split())


def normalize_nico(value):
	text = clean_text(value)
	if not text:
		return ""
	digits = "".join(char for char in text if char.isdigit())
	if len(digits) > 2:
		return ""
	return digits.zfill(2)


def normalize_fraccion(value):
	text = clean_text(value)
	if not text:
		return ""
	return text


def extract_xlsx_links(page_url, timeout):
	response = requests.get(page_url, timeout=timeout)
	response.raise_for_status()

	soup = BeautifulSoup(response.text, "html.parser")
	links = []
	seen = set()

	for anchor in soup.select("a[href]"):
		href = anchor.get("href", "").strip()
		if ".xlsx" not in href.lower():
			continue
		full_url = urljoin(page_url, href)
		if full_url in seen:
			continue
		seen.add(full_url)
		links.append(full_url)

	return links


def iter_catalog_rows(workbook_content):
	workbook = load_workbook(io.BytesIO(workbook_content), data_only=True, read_only=True)

	for sheet_name in workbook.sheetnames:
		sheet = workbook[sheet_name]
		for row in sheet.iter_rows(values_only=True):
			fraccion = normalize_fraccion(row[1] if len(row) > 1 else "")
			nico = normalize_nico(row[2] if len(row) > 2 else "")
			description = clean_text(row[3] if len(row) > 3 else "")
			fraccion_correlativa = normalize_fraccion(row[5] if len(row) > 5 else "")

			if not nico or not description:
				continue

			base_fraccion = ""
			if FULL_FRACCION_RE.match(fraccion):
				base_fraccion = fraccion
			elif FULL_FRACCION_RE.match(fraccion_correlativa):
				base_fraccion = fraccion_correlativa

			if not base_fraccion:
				continue

			full_code = base_fraccion.replace(".", "") + nico
			yield full_code, description


def collect_records(xlsx_urls, timeout, verbose=False):
	records_by_code = {}

	for index, workbook_url in enumerate(xlsx_urls, start=1):
		if verbose:
			print(f"[{index}/{len(xlsx_urls)}] Downloading {workbook_url}")

		response = requests.get(workbook_url, timeout=timeout)
		response.raise_for_status()

		for code, name in iter_catalog_rows(response.content):
			# Keep the first occurrence; repeated codes are common across updates.
			records_by_code.setdefault(code, name)

	records = []
	for code in sorted(records_by_code):
		records.append(
			{
				"id": f"c_fraccion_{code}",
				"code": code,
				"name": records_by_code[code],
				"active": "True",
			}
		)
	return records


def parse_args():
	parser = argparse.ArgumentParser(
		description="Generate Odoo CSV for l10n_mx_catalogs.c_fraccion from SNICE LIGIE + NICO workbooks."
	)
	parser.add_argument("--source-url", default=SOURCE_URL, help="Page URL containing workbook links.")
	parser.add_argument(
		"--output",
		default=DEFAULT_OUTPUT,
		help="Output CSV path. Relative paths are resolved from this script directory.",
	)
	parser.add_argument(
		"--chapter-url",
		action="append",
		default=[],
		help="Optional workbook URL to parse (repeatable). If omitted, URLs are discovered from --source-url.",
	)
	parser.add_argument("--timeout", type=int, default=45, help="HTTP timeout in seconds.")
	parser.add_argument("--verbose", action="store_true", help="Print progress details.")
	return parser.parse_args()


def resolve_output_path(raw_path):
	path = Path(raw_path)
	if path.is_absolute():
		return path
	return Path(__file__).resolve().parent / path


def write_csv(output_path, rows):
	output_path.parent.mkdir(parents=True, exist_ok=True)
	with output_path.open("w", newline="", encoding="utf-8") as csv_file:
		writer = csv.DictWriter(csv_file, fieldnames=["id", "code", "name", "active"])
		writer.writeheader()
		writer.writerows(rows)


def main():
	args = parse_args()
	output_path = resolve_output_path(args.output)

	chapter_urls = args.chapter_url or extract_xlsx_links(args.source_url, args.timeout)
	if not chapter_urls:
		raise SystemExit("No workbook links were found.")

	rows = collect_records(chapter_urls, args.timeout, verbose=args.verbose)
	write_csv(output_path, rows)

	print(f"Generated {len(rows)} rows at {output_path}")


if __name__ == "__main__":
	main()

